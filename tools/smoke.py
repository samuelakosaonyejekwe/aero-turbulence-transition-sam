"""tools/smoke.py - exercise every code path in seconds, not in a quarter hour.

The expensive stages of this project (gen_validation.py, the amplification
database, build_docx.py) take minutes to tens of minutes, and the failures that
matter - a renamed dict key, a changed signature, a NaN escaping a closure, an
array whose length no longer matches its stations - are all detectable in the
first second of the first solve.  Discovering them at minute fourteen has cost
this project whole runs.

This module calls the real functions, at the smallest resolution that still
visits every branch, and asserts the properties that must hold rather than the
numbers, which are the regeneration's job.  It is the gate to run after every
edit and before any full regeneration.

    python3 tools/smoke.py            # all checks
    python3 tools/smoke.py -k kernel  # only checks whose name contains "kernel"

Exit status is the number of failures, so it can gate a commit or a pipeline.
"""
import os
import sys
import time
import traceback

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import utss_paths  # noqa: F401  (anchors ROOT and the working directory)

import matplotlib
matplotlib.use("Agg")

_CHECKS = []
_FAILURES = []


def check(fn):
    _CHECKS.append(fn)
    return fn


def _finite(name, a):
    a = np.asarray(a, dtype=float)
    if not np.all(np.isfinite(a[~np.isnan(a)])):
        raise AssertionError("%s contains inf" % name)


# ----------------------------------------------------------------------
# 1.  imports - the cheapest way to catch a syntax or signature error
# ----------------------------------------------------------------------
@check
def imports_solver():
    """every solver module imports and exposes what its callers use"""
    import case_config as C
    import stability as st
    import utss_solver as U
    for mod, names in ((U, ("march_bl", "solve_airfoil", "solve_flat_plate",
                            "panel_solve", "velocity_field", "CAL")),
                       (st, ("sigma_curve", "twoeq_HL", "H_from_Hstar",
                             "crossflow_factor", "load_database")),
                       (C, ("CRUISE", "CLIMB", "WING", "VALIDATION",
                            "NLF0416", "SWEPT", "SWEPT2"))):
        for n in names:
            assert hasattr(mod, n), "%s lacks %s" % (mod.__name__, n)


@check
def imports_generators():
    """every generator imports without executing its pipeline"""
    for m in ("gen_geometry", "gen_mesh_setup", "gen_postprocessing",
              "gen_equations", "gen_validation", "run_solution",
              "verify_outputs"):
        __import__(m)


# ----------------------------------------------------------------------
# 2.  the stability tables the march reads
# ----------------------------------------------------------------------
@check
def stability_tables():
    """closures are monotone, finite and span the range the march clips to"""
    import stability as st
    Hg, Hs, L, D = st.twoeq_closure()
    assert np.all(np.diff(Hg) > 0), "H grid not increasing"
    assert np.all(np.diff(Hs) < 0), "H* must fall monotonically with H"
    for a, n in ((Hg, "H"), (Hs, "H*"), (L, "l"), (D, "d")):
        _finite(n, a)
    # the inversion the two-equation march depends on
    for h in (2.2, 2.59129, 3.0, 3.5, 3.9):
        assert abs(st.H_from_Hstar(st.twoeq_HL(h)[0]) - h) < 5e-3, \
            "H*(H) inversion off at H=%.3f" % h
    lam, K = st.crossflow_table()
    assert np.all(np.diff(lam) > 0) and np.all(K >= 0.0)


@check
def stability_reverse_branch():
    """the reverse-flow branch is present - the bubble closure depends on it"""
    import stability as st
    eta, prof = st._combined_family()
    Hmax = max(p[0] for p in prof)
    assert Hmax >= st.H_REVERSE - 1e-6, (
        "combined family reaches only H=%.3f but H_REVERSE=%.3f; the reverse "
        "branch failed to build and every bubble length is silently wrong"
        % (Hmax, st.H_REVERSE))
    s = st.sigma_curve(st.H_REVERSE, 400.0)
    assert s.max() > 0.02, "reverse-flow amplification rate collapsed to %.4g" % s.max()


@check
def amplification_database():
    """the tabulated rates are loadable, finite and amplify where they must"""
    import stability as st
    H, R, O, S = st.load_database()
    assert S.shape == (H.size, R.size, O.size)
    _finite("sigma", S)
    assert st.sigma_curve(2.59129, 150.0).max() <= 0.0, \
        "Blasius amplifies below its neutral point"
    assert st.sigma_curve(2.59129, 600.0).max() > 0.0, \
        "Blasius does not amplify well above its neutral point"


# ----------------------------------------------------------------------
# 3.  the inviscid solution
# ----------------------------------------------------------------------
@check
def correlations_against_their_sources():
    """every published correlation, against the form its source publishes"""
    import case_config as C
    import utss_solver as U

    # -- Thwaites, as fitted in White, Viscous Fluid Flow -------------------
    assert abs(U._thwaites_HL(0.0)[1] - 0.220) < 5e-4, "l(0) is not 0.220"
    assert abs(U._thwaites_HL(0.0)[0] - 2.61) < 5e-3, "H(0) is not the fit's 2.61"
    assert abs(U._thwaites_HL(-0.09)[0] - (2.088 + 0.0731/0.05)) < 1e-9
    assert abs(U._thwaites_HL(-0.09)[1]) < 6e-3, "l does not vanish at lambda = -0.09"

    # -- Head (1958) H1(H), and the inverse the march marches on ------------
    assert abs(U._head_H1(1.6) - (3.3 + 0.8234*0.5**-1.287)) < 1e-9
    assert abs(U._head_H1(1.6) - (3.3 + 1.5501*(1.6 - 0.6778)**-3.064)) < 0.03, \
        "the two H1 branches do not meet at H = 1.6"
    for H in (1.3, 1.45, 1.8, 2.4):
        assert abs(U._head_H_from_H1(U._head_H1(H)) - H) < 3e-3, \
            "H1 inversion off at H = %.2f" % H

    # -- Ludwieg & Tillmann (1950) -----------------------------------------
    assert abs(U._ludwieg_tillmann(1.4, 1000.0)
               - 0.246*10**(-0.678*1.4)*1000.0**-0.268) < 1e-12

    # -- Abu-Ghannam & Shaw (1980) -----------------------------------------
    assert abs(U._ags_re_theta_t(3.0, 0.0)
               - (163.0 + np.exp(6.91*(1.0 - 3.0/6.91)))) < 1e-9

    # -- Drela & Giles (1987), the envelope behind the use_os_db ablation ---
    h = 2.59129; hm = 1.0/(h - 1.0)
    assert abs(np.log10(U._re_theta_crit(h))
               - ((1.415*hm - 0.489)*np.tanh(20*hm - 12.9) + 3.295*hm + 0.44)) < 1e-12
    a = 2.4*h - 3.7 + 2.5*np.tanh(1.5*h - 4.65)
    assert abs(U._dn_dReth(h) - 0.01*np.sqrt(a*a + 0.25)) < 1e-12

    # -- Mack (1977) --------------------------------------------------------
    # The FORM is what is published; the anchor is this work's own units
    # conversion and is read off the function rather than duplicated here.
    # Duplicating it meant that re-setting the anchor for the corrected
    # amplification rates failed this check on the constant, not on Mack.
    import inspect
    _anchor = inspect.signature(U._n_crit).parameters["anchor"].default
    assert abs(U._n_crit(1.0) - (-8.43 - 2.4*np.log(0.01) - _anchor)) < 1e-9
    # and Tu is clamped to the interval the correlation is quoted over
    assert U._n_crit(0.02) == U._n_crit(0.08), "the low-Tu clamp is not at 0.08 %"
    assert U._n_crit(5.0) == U._n_crit(3.0), "the high-Tu clamp is not at 2.98 %"

    # -- Karman-Tsien against the EXACT isentropic stagnation pressure ------
    M, g = 0.42, 1.4
    b = np.sqrt(1 - M*M)
    kt = 1.0/(b + (M*M/(1 + b))*0.5)
    exact = 2/(g*M*M)*((1 + (g - 1)/2*M*M)**(g/(g - 1)) - 1)
    assert abs(kt - exact) < 0.006, "Karman-Tsien stagnation Cp %.4f vs exact %.4f" % (kt, exact)
    assert kt < 1.0/b, "Karman-Tsien should stay below Prandtl-Glauert"

    # -- Dhawan & Narasimha, the two forms of one law ----------------------
    assert abs(9.0/0.664**1.5 - 16.63) < 5e-3

    # -- the reverse-flow amplification rates the report and README quote ---
    # 0.0435 at Re_theta = 400 appears in Eq. E12b and in the README, and
    # 0.0417/0.0461 bound it in two solver comments.  They are properties of
    # the committed database and nothing checked that the database still has
    # them.
    import stability as _st
    for Re, want in ((200.0, 0.0417), (400.0, 0.0435), (8000.0, 0.0461)):
        got = float(_st.sigma_curve(_st.H_REVERSE, Re).max())
        assert abs(got - want) < 5e-5, \
            "sigma(H_REVERSE, Re_theta=%g) = %.4f, quoted as %.4f" % (Re, got, want)

    # -- the flight conditions against Sutherland and the ISA --------------
    # These are typed into case_config, and nothing checked that they are the
    # gas they claim to be.
    for nm, d, T in (("cruise", C.CRUISE, 216.65), ("climb", C.CLIMB, 268.66)):
        mu = U._sutherland_mu(T)
        assert abs(d["mu_inf"] - mu) < 1e-8, \
            "%s mu_inf %.4e is not Sutherland at %.2f K (%.4e)" % (nm, d["mu_inf"], T, mu)
        assert abs(d["a_sound"] - np.sqrt(1.4*287.05*T)) < 0.02, \
            "%s speed of sound is not sqrt(gamma R T)" % nm
        assert abs(d["T_inf_K"] - T) < 1e-9
    assert abs(C.CRUISE["rho_inf"] - 22632.0/(287.05*216.65)) < 5e-5, \
        "cruise density is not p/(R T) at the quoted FL360 pressure"

    # -- the reference-temperature closure ---------------------------------
    # exactly 1 at M_e = 0 either way, so no incompressible case can move
    for Te in (None, np.array([216.65])):
        for lam in (True, False):
            assert abs(float(U._ref_temp_nu(np.array([0.0]), laminar=lam,
                                            Te_K=Te)[0]) - 1.0) < 1e-15
    # and Sutherland must exceed the room-temperature power law at 216 K,
    # because omega there is 0.838 and not 0.76
    pw = float(U._ref_temp_nu(np.array([0.42]), laminar=False)[0])
    su = float(U._ref_temp_nu(np.array([0.42]), laminar=False,
                              Te_K=np.array([216.65]))[0])
    assert su > pw, "Sutherland factor %.5f below the power law %.5f" % (su, pw)
    assert abs(su - pw)/pw < 0.01, "the two viscosity laws differ by more than 1 %"


@check
def panel_solution():
    """panel method: closed loop, stagnation pressure, compressibility"""
    import case_config as C
    from utss_solver import panel_solve
    X, Y = C.nlf16_panel_points(60)
    assert abs(X[0] - X[-1]) < 1e-12 and abs(Y[0] - Y[-1]) < 1e-12, "loop not closed"
    xc, yc, Cp, V, th, S = panel_solve(X, Y, 2.0)
    _finite("Cp", Cp); _finite("V", V)
    assert 0.9 < Cp.max() <= 1.0001, "stagnation Cp = %.4f" % Cp.max()
    # incompressible identity, and Karman-Tsien must raise the suction peak
    assert abs(abs(V).max() - np.sqrt(1.0 - Cp.min())) < 1e-9
    Cp_c = panel_solve(X, Y, 2.0, mach=0.42)[2]
    assert Cp_c.min() < Cp.min(), "compressibility did not deepen the suction peak"


@check
def panel_method_invariants():
    """symmetry, antisymmetry and Kutta-Joukowski, which the method must obey exactly"""
    import case_config as C
    import gen_validation as V
    from utss_solver import panel_solve

    def forces(X, Y, a):
        xc, yc, Cp, Vt, th, S = panel_solve(X, Y, a)
        nx = -np.sin(th); ny = np.cos(th)
        Cn = -np.sum(Cp*ny*S); Ca = -np.sum(Cp*nx*S)
        al = np.radians(a)
        return Cn*np.cos(al) - Ca*np.sin(al), Cp, Vt, S

    # The case-study section has a FINITE-ANGLE trailing edge, which is why
    # Squire-Young cannot be evaluated there: the trailing edge is an inviscid
    # stagnation point and the edge velocity goes to zero physically.  If the
    # section is ever changed to a cusp, the reasoning in squire_young's
    # docstring - and the evaluation station it forces - has to be revisited.
    co = C.nlf16_coords(n=2000)
    su = (co["yu"][-1]-co["yu"][-40])/(co["xu"][-1]-co["xu"][-40])
    sl = (co["yl"][-1]-co["yl"][-40])/(co["xl"][-1]-co["xl"][-40])
    te_angle = abs(np.degrees(np.arctan(su) - np.arctan(sl)))
    assert te_angle > 5.0, (
        "the trailing edge is now near-cusped (%.2f deg included); Squire-Young "
        "may be evaluable at it, and squire_young's docstring argues from the "
        "wedge" % te_angle)

    # A SYMMETRIC section at zero incidence must carry exactly no lift, and its
    # two surfaces must see the same pressure.  Nothing tested this, and it is
    # the one case where the answer is known to the last bit.
    X, Y = V._section_points(C.SWEPT2["section"], n=160)   # NACA 64(2)A015
    cl0, Cp0, _, _ = forces(X, Y, 0.0)
    assert abs(cl0) < 1e-12, "symmetric section carries c_l = %.3e at alpha = 0" % cl0

    # and lift must be exactly odd in incidence
    for a in (2.0, 5.0):
        cp, _, _, _ = forces(X, Y, a)
        cm, _, _, _ = forces(X, Y, -a)
        assert abs(cp + cm) < 1e-12, \
            "c_l is not odd in incidence at %.1f deg: %+.6f / %+.6f" % (a, cp, cm)

    # Kutta-Joukowski: the pressure integral and the circulation must agree to
    # the discretisation, and the Kutta condition must equalise the pressure
    # either side of the trailing edge.
    Xn, Yn = C.nlf16_panel_points(130)
    for a in (0.0, 1.5, 4.0):
        cl, Cp, Vt, S = forces(Xn, Yn, a)
        cl_gamma = 2.0*float(np.sum(Vt*S))       # unit chord, unit free stream
        assert abs(cl - cl_gamma) < 2e-3, \
            "lift from pressure %.6f vs from circulation %.6f at %.1f deg" % (cl, cl_gamma, a)
        assert abs(Cp[0] - Cp[-1]) < 5e-3, \
            "Kutta condition: Cp jumps %.2e across the trailing edge" % abs(Cp[0]-Cp[-1])


@check
def section_coordinates_against_their_source():
    """the NLF(1)-0416 coordinate file against Table I of NASA TP-1861"""
    import case_config as C
    # Rows recovered verbatim from the report's own text.  The whole validation
    # rests on the digitised inputs being what their sources say, and nothing
    # checked the one input that CAN be checked against a printed table - the
    # transition locations come off a raster figure and cannot.
    table_I_upper = [
        (0.00049, 0.00403), (0.00509, 0.01446), (0.01393, 0.02573),
        (0.02687, 0.03729), (0.04383, 0.04870), (0.06471, 0.05964),
        (0.08936, 0.06984), (0.11761, 0.07904), (0.14925, 0.08707),
        (0.18404, 0.09374), (0.22169, 0.09892), (0.26187, 0.10247),
        (0.30422, 0.10425), (0.34839, 0.10405), (0.39438, 0.10162),
    ]
    pts = []
    for line in open(C.NLF0416["section"]):
        q = line.split()
        if len(q) == 2:
            try:
                pts.append((float(q[0]), float(q[1])))
            except ValueError:
                pass
    a = np.array(pts)
    up = a[:int(np.argmin(a[:, 0])) + 1][::-1]          # LE -> TE, upper
    for xs, ys in table_I_upper:
        j = int(np.argmin(np.abs(up[:, 0] - xs)))
        assert abs(up[j, 0] - xs) < 2e-5 and abs(up[j, 1] - ys) < 2e-5, \
            ("nlf1_0416.dat departs from TP-1861 Table I at x/c = %.5f: "
             "file has (%.5f, %.5f), the report has (%.5f, %.5f)"
             % (xs, up[j, 0], up[j, 1], xs, ys))
    # and the chord the report states
    assert abs(C.NLF0416["chord_m"] - 0.60902) < 1e-9, \
        "chord is not the 60.902 cm TP-1861 tested at"
    assert abs(C.NLF0416["mach"] - 0.10) < 1e-9, \
        "Fig. 9 of TP-1861 is captioned M = 0.10"
    # no transition data above R = 4e6: the report says the tunnel noise
    # swamped the microphone there
    assert max(C.NLF0416["data"]) <= 4.0e6 + 1, \
        "TP-1861 reports no transition measurements above R = 4.0e6"
    # Boltz: TN D-338 tested -3 to +3 degrees of incidence, 0 to 50 of sweep
    assert all(-3.0 <= a_ <= 3.0 for a_ in C.SWEPT2["alpha_deg"]), \
        "an incidence outside the -3 to +3 deg TN D-338 tested"
    assert all(0.0 <= s_ <= 50.0 for s_ in C.SWEPT2["sweep_deg"]), \
        "a sweep angle outside the 0 to 50 deg TN D-338 tested"

    # Dagenhart & Saric, Table 2 of NASA/TP-1999-209344, verbatim.  This one is
    # a printed table rather than a figure, so it can be checked line by line -
    # and the cross-flow constant is set on it, which makes it the single most
    # load-bearing dataset in the project.
    table_2 = [(1.92e6, 0.78), (2.19e6, 0.73), (2.37e6, 0.58),
               (2.73e6, 0.45), (3.27e6, 0.33), (3.73e6, 0.30)]
    got = list(zip(C.SWEPT["Re_c"], C.SWEPT["x_tr_c"]))
    assert len(got) == len(table_2), \
        "SWEPT has %d points; Table 2 has %d" % (len(got), len(table_2))
    for (r_w, x_w), (r_g, x_g) in zip(table_2, got):
        assert abs(r_g - r_w)/r_w < 1e-6 and abs(x_g - x_w) < 1e-9, \
            ("SWEPT departs from TP-1999-209344 Table 2: file has "
             "(%.3e, %.3f), the report has (%.3e, %.3f)" % (r_g, x_g, r_w, x_w))
    assert abs(C.SWEPT["sweep_deg"] - 45.0) < 1e-9
    assert abs(C.SWEPT["chord_m"] - 1.83) < 1e-9, \
        "the report states a 1.83 m streamwise chord"
    assert abs(C.SWEPT["alpha_deg"] + 4.0) < 1e-9


@check
def stability_against_published_eigenvalues():
    """Orr-Sommerfeld and Falkner-Skan against values published outside this work"""
    import stability as st

    # Falkner-Skan wall shear, the standard table
    for beta, want in ((1.0, 1.23259), (0.5, 0.92768), (0.0, 0.46960),
                       (-0.1, 0.31927), (-0.19, 0.08570)):
        f2 = st.falkner_skan(beta)[2][0]
        assert abs(f2 - want) < 5e-5, \
            "f''(0) at beta=%+.2f is %.6f, published %.5f" % (beta, f2, want)

    # The standard Blasius Orr-Sommerfeld benchmark: at Re_delta* = 998 and
    # alpha delta* = 0.308 the temporal eigenvalue is c = 0.36412 + 0.00796i.
    # This module works in momentum-thickness units, so both are divided by H.
    H = 2.59129
    pr = st.fs_profile_for_H(H)
    U, dU, Upp, y, D1 = st._os_profile(pr, 60, 60.0, 6.0)
    c = st.os_temporal(y, D1, U, Upp, 0.308/H, 998.0/H)
    assert c is not None, "no unstable mode found at the Blasius benchmark point"
    assert abs(c.real - 0.36412) < 2e-4, \
        "phase speed %.5f against the published 0.36412" % c.real
    # The growth rate used to carry a systematic 1.2 per cent deficit, which
    # was fs_profile_for_H blending two similarity profiles instead of solving
    # one: a blend has the requested H and satisfies no similarity equation,
    # and U'' - which drives the instability - was wrong by 4e-4.  Solved, the
    # same code returns the benchmark to better than a tenth of a per cent, so
    # this is held to 0.5 and not to the 3 that tolerated the defect.
    assert abs(c.imag - 0.00796)/0.00796 < 0.005, \
        "growth rate %.6f against the published 0.00796 - more than 0.5 %%" % c.imag


@check
def section_panel_loops_are_closed():
    """every section the project panels is a closed loop, not just the designed one"""
    import case_config as C
    import gen_validation as V
    secs = [("UTSS-NLF16", None),
            ("NLF(2)-0415", V.NLF415),
            ("NACA 64(2)A015", C.SWEPT2["section"]),
            ("NLF(1)-0416", C.NLF0416["section"])]
    for name, fn in secs:
        X, Y = (C.nlf16_panel_points(60) if fn is None
                else V._section_points(fn, n=120))
        gap = float(np.hypot(X[0] - X[-1], Y[0] - Y[-1]))
        assert gap < 1e-12, (
            "%s panels do not close: gap %.6f c.  panel_solve is a closed-body "
            "method and the Kutta condition presumes one trailing-edge point; "
            "the NACA 64(2)A015 is tabulated with a blunt trailing edge and "
            "came back open by 0.00064c." % (name, gap))
        assert len(X) == len(Y) and len(X) > 8


@check
def offbody_field_matches_surface():
    """the off-body field reproduces the surface solution just off the wall"""
    import case_config as C
    from utss_solver import panel_solve, velocity_field
    X, Y = C.nlf16_panel_points(60)
    xc, yc, Cp, V, th, S = panel_solve(X, Y, 1.5)
    d = 0.004
    Px = xc - np.sin(th)*d; Py = yc + np.cos(th)*d
    Cf = velocity_field(X, Y, 1.5, Px, Py, U=1.0)[2]
    e = float(np.mean(np.abs(Cf[8:-8] - Cp[8:-8])))
    assert e < 0.06, "off-body vs surface mean |dCp| = %.4f" % e


# ----------------------------------------------------------------------
# 4.  the boundary-layer march and every transition mechanism
# ----------------------------------------------------------------------
def _plate(**kw):
    from utss_solver import solve_flat_plate
    return solve_flat_plate(kw.pop("L", 1.7), kw.pop("U", 5.4),
                            kw.pop("nu", 1.5e-5), kw.pop("Tu", 3.043),
                            npts=kw.pop("npts", 220), **kw)


@check
def onset_station_is_unambiguous():
    """the state column marks onset where the march fired, and gamma does not"""
    r = _plate()
    i_tr = r["i_tr"]
    assert i_tr is not None
    g = np.asarray(r["gamma"], float)
    st = np.asarray(r["state"], dtype=object)
    # Narasimha's gamma is 1 - exp(-0.412 xi^2) and xi = 0 at onset, so gamma
    # is EXACTLY zero at the station the kernel fired at.  A consumer looking
    # for the first gamma > 0 lands one station downstream, which is how the
    # figures came to mark transition 0.012c aft of the tables.
    assert g[i_tr] == 0.0, "gamma at onset is %g, not 0" % g[i_tr]
    assert st[i_tr] != "laminar", "the onset station is still labelled laminar"
    assert all(s == "laminar" for s in st[:i_tr]), "a station below onset is not laminar"
    first_pos = int(np.argmax(g > 1e-6))
    assert first_pos == i_tr + 1, (
        "gamma first exceeds 1e-6 at %d, onset is %d - consumers must key on "
        "the state column, not on gamma" % (first_pos, i_tr))


@check
def march_output_contract():
    """every consumer key is present, right length, and free of sentinels"""
    r = _plate()
    n = len(r["s"])
    need = ("s Ue theta H Cf Re_theta lam gamma state Re_theta_t n_factor "
            "n_crit n_cf mechanism i_tr i_sep x_tr onset_mech H_lam theta_lam "
            "H_turb theta_turb bubble_burst").split()
    for k in need:
        assert k in r, "march_bl output lost key %r" % k
    for k in ("s Ue theta H Cf Re_theta lam gamma Re_theta_t n_factor n_crit "
              "n_cf state mechanism H_lam theta_lam H_turb theta_turb").split():
        assert len(r[k]) == n, "%s has length %d, stations %d" % (k, len(r[k]), n)
    for k in ("theta", "H", "Cf", "Re_theta", "gamma", "n_factor", "n_crit"):
        _finite(k, r[k])
    rt = np.asarray(r["Re_theta_t"], float)
    assert not np.any(rt[~np.isnan(rt)] >= 1e8), \
        "the 1e9 'branch inactive' sentinel escaped into Re_theta_t"
    assert np.all(np.asarray(r["gamma"]) >= 0.0) and np.all(np.asarray(r["gamma"]) <= 1.0)
    assert np.all(np.asarray(r["H"])[:r["i_tr"] or n] > 1.0), "shape factor below unity"


@check
def kernel_every_mechanism_fires():
    """each of the four mechanisms is reachable and labels itself"""
    import case_config as C
    from utss_solver import solve_airfoil
    got = {}
    # bypass: high free-stream turbulence on a plate
    got["bypass"] = _plate(Tu=3.043, L_turb=1.53e-3)["onset_mech"]
    # natural / TS: a quiet plate
    got["TS-natural"] = _plate(L=4.0, U=27.0, Tu=0.03, npts=320)["onset_mech"]
    # separation: the T3C4 rig geometry, adverse gradient at low speed
    got["separation"] = _plate(L=1.5, U=1.51, Tu=2.11, npts=320,
                               dUe=-0.35)["onset_mech"]
    # cross-flow: a swept section
    X, Y = C.nlf16_panel_points(80)
    r = solve_airfoil(X, Y, -2.0, 40.0, 1.5e-5, 1.83, 0.02, sweep_deg=45.0)
    got["crossflow"] = r["surfaces"]["upper"]["onset_mech"]
    for want, seen in got.items():
        assert seen == want, "expected %s, kernel selected %r" % (want, seen)


@check
def kernel_continuous_in_turbulence():
    """the answer is continuous in Tu across the natural/bypass handover"""
    import case_config as C
    from utss_solver import solve_airfoil
    cr, W = C.CRUISE, C.WING
    npan = 80
    X, Y = C.nlf16_panel_points(npan)

    def sweep(n, lo=0.08, hi=0.30):
        tus = np.linspace(lo, hi, n)
        out = []
        for tu in tus:
            r = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                              W["MAC"], float(tu), sweep_deg=W["le_sweep_deg"],
                              mach=cr["mach"])
            u = r["surfaces"]["upper"]["x_tr_chord"]
            out.append(1.0 if u != u else float(u))
        return tus, np.array(out)

    # transition may only move forward as the stream gets noisier
    tus, xtr = sweep(23)
    assert np.all(np.diff(xtr) <= 1e-9), \
        "transition moves aft with rising Tu: %s" % np.array2string(xtr, precision=3)

    # Continuity is tested by REFINEMENT, not by an absolute threshold.  The
    # predicted transition location can only land on a panel, so the sampled
    # jump can never fall below the panel spacing however smooth the model is;
    # an absolute bound therefore measures the grid, not the physics.  Halving
    # the sampling halves the jump for a continuous function and leaves it
    # unchanged for a step.  The old gate produced a 0.17c step here that did
    # not refine away at all.
    j1 = float(np.max(np.abs(np.diff(sweep(12)[1]))))
    j2 = float(np.max(np.abs(np.diff(xtr))))
    xs = np.sort(np.abs(np.asarray(C.nlf16_coords(npan)["xu"])))
    spacing = float(np.max(np.diff(xs)[(xs[:-1] > 0.2) & (xs[:-1] < 0.6)]))
    assert j2 <= max(0.55*j1, 1.6*spacing) + 1e-9, (
        "the jump in transition location did not refine away: %.4fc at "
        "dTu = %.4f, %.4fc at half that, against a panel spacing of %.4fc. "
        "A step that survives refinement is a discontinuity in the model."
        % (j1, (tus[1]-tus[0])*2, j2, spacing))


@check
def bubble_closure():
    """a separating plate forms a bubble, marches it, and closes it"""
    r = _plate(L=1.5, U=1.51, Tu=2.11, npts=320, dUe=-0.35)
    assert r["i_sep"] is not None, "no separation on a strongly decelerated plate"
    assert r["onset_mech"] == "separation"
    i0, i1 = r["i_sep"], r["i_tr"]
    assert i1 > i0, "bubble closed before it opened"
    th = np.asarray(r["theta"], float)
    assert th[i1] >= th[i0]*0.99, "momentum thickness fell across the dead air"
    assert np.all(np.asarray(r["Cf"])[i0:i1] <= 1e-12), "wall shear inside the bubble"


@check
def bubble_can_reattach_laminar():
    """a bubble in a gradient that recovers reattaches instead of absorbing"""
    from utss_solver import solve_flat_plate
    # decelerate to separate, then accelerate: the classic short bubble in a
    # gradient that relaxes.  Before laminar reattachment existed the layer
    # could only transition or be declared burst at the trailing edge.
    n = 600
    s = np.linspace(1e-4, 1.5, n)
    Ue = 1.9 - 0.9*np.exp(-((s - 0.75)/0.16)**2)     # dip and recovery
    from utss_solver import march_bl
    r = march_bl(s, Ue, 1.5e-5, Tu_pct=0.05)
    # i_sep is the sticky record of the first separation and must survive the
    # reattachment that clears the live bubble state
    assert r["i_sep"] is not None, "no separation in the decelerating region"
    assert not r["bubble_burst"], "reattached, yet still reported as burst"
    assert r.get("n_reattach", 0) >= 1, (
        "the bubble never reattached even though the gradient recovers; the "
        "separated state is still absorbing")
    # the amplification accumulated across the dead air must not be thrown away
    nf = np.asarray(r["n_factor"], float)
    i0 = r["i_sep"]
    assert np.nanmax(nf[i0:]) > 0.0, "amplification vanished at reattachment"
    tail = nf[min(i0 + 40, n - 1):]
    assert tail.size and np.nanmax(tail) >= 0.9*np.nanmax(nf[:i0 + 40]), (
        "the amplification factor collapsed after reattachment: the bubble's "
        "own N was discarded instead of being carried into the attached "
        "integral")


@check
def compressible_closures():
    """raising the Mach number delays transition and thins nothing to zero"""
    import case_config as C
    from utss_solver import solve_airfoil
    cr, W = C.CRUISE, C.WING
    X, Y = C.nlf16_panel_points(80)
    out = {}
    for M in (0.0, 0.42):
        r = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                          W["MAC"], cr["Tu_pct"], sweep_deg=W["le_sweep_deg"],
                          mach=M)
        out[M] = r
        _finite("Cd", [r["Cd"]]); _finite("Cl", [r["Cl"]])
        assert r["Cd"] > 0.0
    assert out[0.42]["Cl"] > out[0.0]["Cl"], "compressibility did not raise the lift"


@check
def swept_drag_factor():
    """the swept drag conversion, against the yawed flat plate it must reproduce"""
    from utss_solver import _swept_drag_factor as F
    # 1. zero sweep is Squire-Young untouched
    for r_ in (0.7, 0.84, 1.0):
        for H in (1.4, 2.15, 2.8):
            assert abs(F(r_, H, 0.0) - r_**((H + 5.0)/2.0)) < 1e-12
            # and sweep_transform=False must recover it at any sweep
            assert abs(F(r_, H, 45.0, swept=False) - r_**((H + 5.0)/2.0)) < 1e-12

    # 2. THE CHECK THAT SETS THE FORM.  A flat plate at yaw is a flat plate in
    #    the free stream, so its drag is the unyawed value at the streamwise run
    #    length: with U_e,n = U_n the factor must be exactly cos(L), for every
    #    sweep and every shape factor.  2 theta_n/c_n * cos(L) = 2 theta_n/c,
    #    and the independence principle gives theta_n = theta_streamwise.
    #    The cos^2(L) this replaced fails here by a whole cos(L) - 29 % at 45
    #    degrees - and so would cos^3(L) on Squire-Young alone, by cos^2(L).
    for L in (0.0, 5.0, 12.0, 30.0, 45.0, 60.0):
        for H in (1.4, 2.15, 2.8, 4.0):
            got = F(1.0, H, L)
            assert abs(got - np.cos(np.radians(L))) < 1e-12, \
                "yawed flat plate at L=%g, H=%g: factor %.6f, cos(L) %.6f" \
                % (L, H, got, np.cos(np.radians(L)))

    # 3. continuous at zero sweep, bounded, and falling once the geometry
    #    dominates.  It is NOT monotone from zero: the factor is cos(L) times a
    #    cos^2/sin^2-weighted mean of r^p and r, and with r < 1 < p the
    #    span-wise term r is the LARGER of the two, so a little sweep raises the
    #    drag by a quarter of a per cent before cos(L) takes over near 20 deg.
    r_, H = 0.84, 2.2
    p_ = (H + 5.0)/2.0
    vals = [F(r_, H, L) for L in (0.0, 1.0, 12.0, 30.0, 45.0)]
    assert abs(vals[0] - r_**p_) < 1e-12, "zero sweep is not Squire-Young"
    assert abs(vals[0] - vals[1]) < 2e-3, "discontinuous at zero sweep"
    for L in (0.0, 5.0, 12.0, 30.0, 45.0, 60.0):
        c = np.cos(np.radians(L))
        lo, hi = c*min(r_**p_, r_), c*max(r_**p_, r_)
        assert lo - 1e-12 <= F(r_, H, L) <= hi + 1e-12, \
            "factor at L=%g is not cos(L) times a mean of r^p and r" % L
    tail = [F(r_, H, L) for L in (20.0, 30.0, 45.0, 60.0)]
    assert all(a > b for a, b in zip(tail, tail[1:])), \
        "drag factor does not fall with sweep beyond 20 deg: %s" % tail

    # 4. the span-wise term is additive, so the factor must exceed the
    #    chordwise-only cos^3(L) Squire-Young term at any real sweep
    for L in (12.0, 45.0):
        r_, H = 0.84, 2.2
        chord_only = np.cos(np.radians(L))**3 * r_**((H + 5.0)/2.0)
        assert F(r_, H, L) > chord_only, "span-wise friction term is missing"


@check
def swept_drag_is_near_unswept_at_small_sweep():
    """12 deg of sweep must not move the viscous drag by five per cent"""
    import case_config as C
    from utss_solver import solve_airfoil
    cr, W = C.CRUISE, C.WING
    X, Y = C.nlf16_panel_points(80)
    cd = {}
    for sw in (0.0, W["le_sweep_deg"]):
        r = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                          W["MAC"], cr["Tu_pct"], sweep_deg=sw, mach=cr["mach"])
        cd[sw] = r["Cd"]
    d = abs(cd[W["le_sweep_deg"]] - cd[0.0])/cd[0.0]
    # A mild sweep leaves the wetted area and the streamwise run length very
    # nearly alone, so the streamwise profile drag has to be very nearly the
    # unswept one.  The cos^2(L) conversion this replaced put it 4.4 % below.
    assert d < 0.02, ("12 deg of sweep moved the section drag by %.1f %%: "
                      "%.2f counts unswept, %.2f swept"
                      % (100*d, cd[0.0]*1e4, cd[W["le_sweep_deg"]]*1e4))


@check
def squire_young_station_is_declared():
    """the drag says where it was evaluated, and whether H there is a clamp"""
    import case_config as C
    from utss_solver import solve_airfoil
    cr, W = C.CRUISE, C.WING
    X, Y = C.nlf16_panel_points(80)
    kw = dict(sweep_deg=W["le_sweep_deg"], mach=cr["mach"],
              T_inf_K=cr["T_inf_K"])

    # the station is a declared constant, not a literal buried in a closure
    from utss_solver import CAL
    assert "sy_x_ref" in CAL and 0.5 < CAL["sy_x_ref"] < 1.0

    # and it moves the drag, so it must be sweepable
    cds = {}
    for xr in (0.90, 0.98):
        r = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                          W["MAC"], cr["Tu_pct"], cal=dict(sy_x_ref=xr), **kw)
        cds[xr] = r["Cd"]*1e4
        u = r["surfaces"]["upper"]
        assert abs(u["x_squire_young"] - xr) < 0.03, \
            "Squire-Young evaluated at %.4f, asked for %.2f" % (u["x_squire_young"], xr)
    assert cds[0.98] > cds[0.90], "the drag no longer rises towards the trailing edge"

    # The rise between stations is FRICTION, not scatter: the drag counted so
    # far plus the friction still ahead is the same number wherever it is
    # evaluated.  This is what says the station is converged rather than
    # uncertain, and it is the statement this project first got backwards.
    _tz = getattr(np, "trapezoid", None) or np.trapz
    tot = {}
    for xr in (0.90, 0.98):
        r = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                          W["MAC"], cr["Tu_pct"], cal=dict(sy_x_ref=xr), **kw)
        cosL = r["cos_sweep"]; chord = W["MAC"]*cosL; Un = cr["U_inf"]*cosL
        omit = 0.0
        for sf in ("upper", "lower"):
            s_ = r["surfaces"][sf]
            x_ = np.asarray(s_["x"], float); cf_ = np.asarray(s_["Cf"], float)
            ue_ = np.asarray(s_["Ue"], float)/Un
            a_ = np.asarray(s_["s"], float)/chord
            m_ = x_ >= xr
            if m_.sum() > 1:
                omit += _tz(cf_[m_]*ue_[m_]**2, a_[m_])
        tot[xr] = r["Cd"]*1e4 + omit*1e4
    assert abs(tot[0.98] - tot[0.90]) < 1.5, (
        "drag-plus-omitted-friction is not station-invariant: %.2f at 0.90c "
        "against %.2f at 0.98c.  The spread across stations would then be a "
        "real uncertainty rather than friction being included."
        % (tot[0.90], tot[0.98]))
    assert abs(cds[0.98] - cds[0.90]) > 1.0, (
        "the evaluation station is worth less than a count now (%.2f vs %.2f); "
        "the sensitivity table and the text that quotes it need re-checking"
        % (cds[0.90], cds[0.98]))

    # Head's clamp: the climb case sits on it at the evaluation station, and
    # the flag has to say so - the drag there is formed from a bound
    cl = C.CLIMB
    r = solve_airfoil(X, Y, cl["alpha_deg"], cl["U_inf"], cl["nu_inf"], W["MAC"],
                      cl["Tu_pct"], sweep_deg=W["le_sweep_deg"],
                      mach=cl["mach"], T_inf_K=cl["T_inf_K"])
    u = r["surfaces"]["upper"]
    assert u["H_sy_at_clip"] == (u["H_te_squire_young"] >= 2.8 - 1e-6), \
        "H_sy_at_clip disagrees with the shape factor it describes"


@check
def airfoil_output_contract():
    """solve_airfoil returns what run_solution and gen_validation read"""
    import case_config as C
    from utss_solver import solve_airfoil
    X, Y = C.nlf16_panel_points(80)
    r = solve_airfoil(X, Y, 1.5, 120.0, 4.0e-5, 1.9, 0.07,
                      sweep_deg=12.0, mach=0.42)
    for k in ("panel", "surfaces", "Cl", "Cd", "alpha", "theta_te_c",
              "alpha_normal_deg", "cos_sweep", "mach_solve", "sweep_transform"):
        assert k in r
    # "alpha" is the incidence the caller asked for and "alpha_normal_deg" the
    # one the panel method was run at; on a swept section they must differ, or
    # the normal-plane transformation has overwritten the streamwise value again
    assert abs(r["alpha"] - 1.5) < 1e-12, "alpha is not the requested incidence"
    assert r["alpha_normal_deg"] > r["alpha"], "normal-plane incidence not raised"
    assert abs(r["mach_solve"] - 0.42*r["cos_sweep"]) < 1e-12
    for surf in ("upper", "lower"):
        s = r["surfaces"][surf]
        for k in ("x", "y", "Cp", "Re_x", "x_tr_chord", "x_sep_chord",
                  "bubble_burst", "theta_te_c", "H_te", "H_te_at_clip",
                  "sep_margin_H", "x_sep_turb_chord", "Me",
                  "Ue_te_ratio", "H_te_squire_young", "x_squire_young",
                  "H_sy_at_clip"):
            assert k in s, "surface dict lost %r" % k
        assert len(s["x"]) == len(s["s"]) == len(s["Cf"]) == len(s["Me"])
        # the edge Mach number must come off the corrected pressure, not off
        # U_e/a_inf; the two differ by over a per cent at cruise
        me = np.asarray(s["Me"], float)
        _finite("Me", me)
        assert me.min() >= 0.0 and me.max() < 1.0, "edge Mach outside [0,1)"


# ----------------------------------------------------------------------
# 5.  the drivers, at the smallest sweep that still writes every column
# ----------------------------------------------------------------------
@check
def wall_normal_profiles():
    """the profile reconstruction runs and returns a monotone, bounded profile"""
    import run_solution as R
    import case_config as C
    from utss_solver import solve_airfoil
    X, Y = C.nlf16_panel_points(80)
    cr, W = C.CRUISE, C.WING
    rc = solve_airfoil(X, Y, cr["alpha_deg"], cr["U_inf"], cr["nu_inf"],
                       W["MAC"], cr["Tu_pct"], sweep_deg=W["le_sweep_deg"],
                       mach=cr["mach"])
    df = R.bl_profiles(rc, write=False)
    assert len(df) > 0
    for st, g in df.groupby("station"):
        u = g.u_Ue.to_numpy(float)
        assert u.min() >= -1e-9 and u.max() <= 1.0 + 1e-9, \
            "%s: u/Ue outside [0,1]" % st
        assert np.all(np.diff(u) >= -1e-6), "%s: velocity profile not monotone" % st
        assert float(g.delta_mm.iloc[0]) > 0.0
        assert float(g.recovery_r.iloc[0]) > 0.0
    # the laminar recovery factor must differ from the turbulent one
    rr = df.recovery_r.to_numpy(float)
    assert rr.min() < rr.max() + 1e-12


@check
def lifting_line_closed_form():
    """the elliptic planform check the pipeline asserts before writing"""
    import run_solution as R
    assert R._lifting_line_check() is True
    ll = R.lifting_line()
    for k in ("CL", "CDi", "e", "a0", "alpha_L0", "sweep_c4"):
        assert k in ll and np.isfinite(ll[k])
    assert 0.0 < ll["e"] <= 1.0


@check
def verify_outputs_contract():
    """verify_outputs can still find every CSV and column it reads"""
    import verify_outputs as V
    want, present, absent = V.checks()
    assert len(want) > 15
    for label, value, anchor in want:
        assert isinstance(value, str) and value, "%s has no value" % label
        assert anchor is None or isinstance(anchor, str)
    # the matcher itself: a number must not pass on a longer one containing it
    txt = V.flatten("Table 14. NLF vs fully-turbulent drag. 1683 and 0.1685")
    assert not V.find_value(txt, "168", "NLF vs fully-turbulent drag")[0], \
        "168 matched inside 1683 - the substring bug is back"
    assert V.find_value(V.flatten("NLF vs fully-turbulent drag. 168.0"),
                        "168", "NLF vs fully-turbulent drag")[0], \
        "168 did not match 168.0"
    # A .docx must be flattened IN DOCUMENT ORDER.  Reading every paragraph and
    # then every table puts each caption tens of thousands of characters from
    # the table it labels, and every anchored check then fails with "present,
    # but not within 2500 chars" on a document whose numbers are all correct.
    import tempfile
    from docx import Document as _D
    d = _D()
    d.add_paragraph("ANCHOR CAPTION HERE")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "value"; t.rows[0].cells[1].text = "42.0"
    d.add_paragraph("TRAILING PARAGRAPH")
    with tempfile.TemporaryDirectory() as td:
        f = os.path.join(td, "order.docx")
        d.save(f)
        txt = V.flatten(V.document_text(f)[0])
    assert txt.index("ANCHOR CAPTION") < txt.index("42.0") < txt.index("TRAILING"), \
        "document_text does not preserve .docx body order: %r" % txt
    assert V.find_value(txt, "42.0", "ANCHOR CAPTION HERE")[0], \
        "a table value is no longer anchorable to the caption above it"

    # Caption numbering must count FIRST occurrences: a repeat is a
    # cross-reference, and the renderer wraps lines wherever it likes, so
    # "... in Table 12. The drag ..." can begin a line and look like a caption.
    good = "Table 1. A.\nTable 2. B.\nsee\nTable 1. And so on.\nTable 3. C."
    bad = "Table 1. A.\nTable 3. C.\nTable 2. B."
    for lab, raw_, want in (("cross-reference", good, True),
                            ("genuinely out of order", bad, False)):
        res = dict((n, o) for n, o, _ in
                   V.structural_checks(V.flatten(raw_), raw_))
        got = res["Table numbering (%d captions)"
                  % (3 if want else 3)] if False else None
        got = [o for n, o, _ in V.structural_checks(V.flatten(raw_), raw_)
               if n.startswith("Table numbering")][0]
        assert got is want, "%s: numbering check returned %s" % (lab, got)


def main():
    only = None
    if "-k" in sys.argv:
        only = sys.argv[sys.argv.index("-k") + 1]
    t0 = time.time()
    run = 0
    for fn in _CHECKS:
        if only and only not in fn.__name__:
            continue
        run += 1
        t = time.time()
        try:
            fn()
            print("  ok    %-34s %6.2f s   %s"
                  % (fn.__name__, time.time() - t, fn.__doc__ or ""))
        except Exception as e:                      # noqa: BLE001
            _FAILURES.append((fn.__name__, e, traceback.format_exc()))
            print("  FAIL  %-34s %6.2f s   %s"
                  % (fn.__name__, time.time() - t, e))
    print("\n%d check(s), %d failed, %.1f s total"
          % (run, len(_FAILURES), time.time() - t0))
    for name, _, tb in _FAILURES:
        print("\n--- %s ---\n%s" % (name, tb))
    return len(_FAILURES)


if __name__ == "__main__":
    sys.exit(main())
